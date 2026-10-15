from pathlib import Path
import re, json, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
report = (ROOT / 'docs/DENIS_FINAL_APP_REVIEW.md').read_text(encoding='utf-8')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, BLUE, TEAL, GREY = '132E48', '1979CE', '24AEB9', '547089'

def box(sl, x,y,w,h,text,size=23,color=NAVY,bold=False):
    tf=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Aptos'; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12)
    return tf

def slide(title,body,tag='CURRENT IMPLEMENTATION',notes='',size=23):
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb=RGBColor.from_string('F5F8FC')
    box(sl,.55,.25,12,.35,tag,12,BLUE,True)
    box(sl,.55,.85,12.2,.9,title,30,NAVY,True)
    sh=sl.shapes.add_shape(1, Inches(.55), Inches(1.8), Inches(12.2), Inches(.035))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(TEAL); sh.line.fill.background()
    box(sl,.65,2.12,12,4.85,body,size)
    box(sl,.55,7.13,11,.2,'NaReT review • 8 October 2026 • Baseline 9d880a1 / current working tree',10,GREY)
    box(sl,12,7.1,.7,.25,str(len(prs.slides)),11,GREY)
    sl.notes_slide.notes_text_frame.text=notes or body
    return sl

slide('NaReT: what Denis built and what changed',
      'A practical comparison of the original app and the merged headset app.\nIncludes scenario rules, speech, CCBM, tracking, CPR feedback, patient appearance and menus.\nOriginal Marek dialogue: all 14 fixed lines, in German and Polish, with English translations.',
      'REVIEW AND DISCUSSION',report,size=25)
slide('Credit the existing work correctly',
      'Denis built the Godot/OpenXR application, scenario engine, patient and witnesses, audio/cloud framework, menus and feedback infrastructure.\nHis setup already imported your hand-zone, compression detector and pose stabilizer. He also integrated an ArUco addon and settling placement.\nThe merge adds your complete three-marker/startup/motion-gate path, restores your CPR display, connects Paul’s detector, and adds English and CCBM.')
slide('Same state machine; a different way to submit actions',
      'Before: progression buttons submitted actions to Denis’s scenario controller.\nNow: speech → recognised text → observation → CCBM decision → Denis’s allowed-action guard → scenario update.\nThe state machine still decides what is legal. The language model cannot freely move the exercise to any stage.',notes='scripts/scenario/scenario_controller.gd; scripts/scenario/ccbm_bridge.gd; scripts/main.gd')
slide('Exactly what the two switches do',
      'cpr_witness_only: exclude AED actions and treatment performed by a virtual bystander. Anna stays; Marek and the AED scene instances are removed.\nobservation_driven: hide routine progression buttons, such as responsiveness, breathing, ventilation and handover.\nStart, reset, calibration and audio controls remain. Physical compression events still come from your detector.',size=24)
slide('Marek: what his original role was',
      'Denis’s original second witness could call 112, fetch an AED, relay dispatcher advice and offer/help with compressions.\nHis 14 fixed lines were triggered by events and timers; they were not a conversation played from start to finish.\nThe current one-witness exercise does not instantiate him. His original source remains in the repository. The full script is included later.',notes='Baseline data/bystander_lines.json and scripts/scenario/bystander_director.gd at 9d880a1.')
slide('Local recognition: what was actually added',
      'Quest microphone → Denis’s existing WAV capture → uni_api_client.gd local ASR route.\nUSB forwarding sends localhost:8767/inference to the PC’s whisper.cpp server in Ubuntu/WSL. It runs the medium model on the PC.\nText returns to the app. Cloud LLM/TTS support remains separate. Local ASR is not bundled inside the APK.',notes='tools/setup_local_asr.sh; tools/start_local_asr.sh; tools/connect_local_speech.py; scripts/api/uni_api_client.gd')
slide('Recognition and interpretation are separate',
      'ASR answers: “What words were spoken?”\nThe observation mapper/interpreter answers: “Was this a completed finding, a request, a question, or just conversation?”\nCCBM evaluates the observation history; Denis’s scenario guard checks the resulting action. Saying “I gave a breath” remains a reported action, not a sensor measurement.')
slide('Denis did more than ArUco detection',
      'Original: an ArUco adapter, configurable marker, tracking status, temporal filter and model calibration. His spatial anchor settled for 1.5 seconds and averaged a stable window.\nAdded: calibrated three-marker common-anchor fusion, robust startup median/rotation medoid, viewer-motion gating and restored original filter settings.\nPaul: newer native detector/addon. Your camera adapter and downstream placement remain in the merged path.')
slide('The complete current placement path',
      'Camera colour conversion → Paul’s detector → capture-time room pose\n→ marker_layout.gd: each marker reconstructs the same anchor\n→ startup_placement.gd: robust initial reference\n→ viewer_motion_gate.gd → original simple_pose_stabilizer.gd\n→ spatial_anchor.gd: saved alignment → Denis’s patient model',size=25)
slide('Transparency is implemented in a shader',
      'mannequin_controller.gd installs patient_overlay.gdshader on the patient materials.\nA 4 × 4 pixel pattern discards a fraction of body pixels. Remaining pixels keep normal opaque depth; this avoids showing multiple transparent layers through each other.\nThe head region stays solid to avoid teeth/eyeballs showing through facial skin. The UI converts transparency to opacity: opacity = 1 − transparency / 100.',notes='assets/patient_overlay.gdshader; scripts/mannequin/mannequin_controller.gd. Screen-space coverage can show a visible pattern in XR.')
slide('Alignment, floor lock and fitting are three different things',
      'Alignment: patient world pose = marker anchor × saved model offset. Controller nudges edit the relative position/rotation.\nFloor orientation: discard marker pitch/roll. Floor height/contact: separate reference-height and posed-model geometry handling. One does not guarantee the other.\nOptional fitting: target head-to-waist length and shoulder width can scale the model. Defaults are zero (disabled); scene scale 0.9 does not prove a measured torso match.')
slide('Head animation: the current values differ from Denis’s',
      'Denis: neck up to 28°; head up to 20°. Current merged code: 10° + 10°. The reduction was a prototype choice, not evidence that Denis’s animation was wrong.\nRecommendation: restore 28°/20° after checking the rig’s direction, while keeping the deliberate two-hand gesture trigger.\nThese are local bone rotations. They do not prove the physical mannequin’s airway was opened.', 'PROPOSED CHANGE — NOT APPLIED BY THIS PRESENTATION')
slide('Menu audit: where the repetition occurs',
      'The old tiny language and microphone controls are hidden. The large language/microphone controls are active.\nStart can still appear both as the large Start control and in the dynamic action list. Training controls also repeat in the calibration panel.\nChoosing a language currently also requests Own-language mode; this coupling makes the behavior harder to understand.',notes='scripts/ui/hand_menu.gd; scripts/ui/calibration_panel.gd; scripts/main.gd')
slide('One Language umbrella; one training control',
      'Start training / Training active\nLanguage: English  →  English • Deutsch • Polski\nMicrophone: On / Off\nAlign patient  →  position, rotation, Save\nMore  →  Test voice, training mode, results/export\nKeep the chosen option green. Flash momentary actions briefly. Use the same selection state in the controller-accessible menu.', 'PROPOSED MENU — NOT INSTALLED',size=25)

# Preserve every review paragraph, with editable text and source notes.
def clean(s):
    s=re.sub(r'\[([^\]]+)\]\([^\)]+\)',r'\1',s)
    return s.replace('**','').replace('`','').strip()
content=report.split('## File inventory')[0]
sections=re.split(r'(?m)^## ',content)[1:]
for sec in sections:
    title,_,rest=sec.partition('\n')
    current=title
    chunks=[]
    for para in re.split(r'\n\s*\n',rest.strip()):
        if para.startswith('### '):
            current=clean(para[4:].replace('\n',' ')); continue
        if not para.strip(): continue
        text=clean(para.replace('\n',' '))
        words=text.split()
        # Max 125 words per detail slide keeps the original audit fully readable.
        for off in range(0,len(words),125):
            chunks.append((current,' '.join(words[off:off+125])))
    for heading,paragraph in chunks:
        slide(heading,paragraph,'DETAILED FILE REVIEW',notes=paragraph,size=23)

raw=subprocess.run(['git','show','9d880a1:data/bystander_lines.json'],cwd=ROOT,capture_output=True,check=True).stdout
lines=json.loads(raw)
translations=[
('I’m here. What should I do?','After accepted responsiveness assessment, with a 2.5-second delay.'),
('I’m calling 112 and putting it on speakerphone.','Accepted emergency-call action.'),
('The dispatcher is on the line. They’re asking whether he is breathing normally.','Five seconds after the emergency-call action.'),
('Should I call the emergency number?','CALL_HELP: no call after 15 seconds.'),
('I’ll get the defibrillator!','Accepted AED request.'),
('Should I look for a defibrillator?','During CPR, call made but no AED requested: more than 20 seconds since the call.'),
('Here is the AED! Turn it on.','Return with AED after the fetch sequence; travel affects timing.'),
('The dispatcher says: press faster!','Four consecutive rate samples below 98 BPM after warm-up; call/phone available; cooldown.'),
('A little slower, says the dispatcher.','Four consecutive rate samples above 124 BPM after warm-up; call/phone available; cooldown.'),
('Hands in the centre of the chest, says the dispatcher.','Three consecutive off-target strokes; call/phone available; cooldown.'),
('The device says: resume compressions immediately!','Post-shock inactivity longer than five seconds.'),
('I can take over compressions if you can’t continue.','Analysis-due event during CPR; this line offers help rather than accepting a swap.'),
('The ambulance crew is here.','Handover, or eight seconds after simulated recovery/ROSC.'),
('I’ll keep doing compressions—attach the pads!','AED handoff/helper-compression sequence.')]
marek=[(k,v) for k,v in lines.items() if isinstance(v,dict) and v.get('speaker')=='marek']
assert len(marek)==len(translations)==14
md=['# Marek: complete original fixed dialogue','', 'Source: Denis baseline commit 9d880a1, data/bystander_lines.json. German and Polish are verbatim; English is a translation prepared for this review. Event-triggered script, not a single linear conversation. Marek is inactive in the current one-witness app.','']
for i,((key,v),(en,trigger)) in enumerate(zip(marek,translations),1):
    body=f"German: {v['de']}\nPolish: {v['pl']}\nEnglish translation: {en}\nWhen: {trigger}"
    slide(f'Marek {i}/14 • {key}',body,'ORIGINAL SCRIPT — INACTIVE IN CURRENT APP',notes=f"Baseline: data/bystander_lines.json; style: {v['style']}. Trigger: bystander_director.gd",size=22)
    md += [f'## {i}. {key}',f"- German: {v['de']}",f"- Polish: {v['pl']}",f'- English translation: {en}',f'- Trigger: {trigger}',f"- Original voice style: {v['style']}",'']
(OUT/'MAREK_FULL_SCRIPT_DE_PL_EN.md').write_text('\n'.join(md),encoding='utf-8')
inventory=report.split('<!-- generated runtime inventory -->')[1]
rows=[clean(x) for x in inventory.splitlines() if x.startswith('| [')]
for i in range(0,len(rows),9):
    slide('Runtime file inventory', '\n'.join(row.strip('| ').replace(' | ','  •  ') for row in rows[i:i+9]),'51 FILES: 17 CHANGED • 13 ADDED • 21 UNCHANGED',size=18)
slide('What to test before calling the app finished',
      'Placement: stationary mannequin, head motion, marker loss, restart and recenter; verify saved nudges and floor height.\nExercise: full English run with natural multi-sentence speech, interruptions, 30 compressions, two reported breaths and handover; then German/Polish.\nDeployment: PC/USB ASR and CCBM remain development dependencies. Connect real mannequin sensors if verified depth/recoil/ventilation are required.', 'VALIDATION STILL REQUIRED')
slide('Sources and scope of this review',
      'Denis baseline: NaReT-Nachbarsprache-wird-reanimiert/naret-xr-app, commit 9d880a1 (28 September 2026).\nCurrent: local working tree, installed version 42 / 0.3.40-original-stabilizer. Paul addon: 729f8c0.\nFull source audit: DENIS_FINAL_APP_REVIEW.md. Reported validation: 24 application suites passed with the blocked Windows native detector excluded. Hardware behavior still needs the checks listed here.',notes=report)
dest=OUT/'NaReT_Denis_vs_Final_Review_20261008.pptx'
prs.save(dest)
print(f'{dest}\nSlides: {len(prs.slides)}')
